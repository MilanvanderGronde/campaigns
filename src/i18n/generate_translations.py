# -*- coding: utf-8 -*-
from pathlib import Path
import openpyxl
from openpyxl.styles import Font
from rsa_translations import LANGS

REPO_ROOT = Path(__file__).resolve().parents[2]
SRC = REPO_ROOT / 'data' / 'reference' / 'Stock_Campaign_RSA_Translations_source.xlsx'
OUT = REPO_ROOT / 'data' / 'reference' / 'Stock_Campaign_RSA_Translations.xlsx'

BUY_WORDS = ['buy','comprar','kaufen','acheter','kupować','kupowac','comprare']

def fits(tmpl, name, limit):
    s = tmpl.format(name=name)
    return s if len(s) <= limit else None

def resolve_h5(t, theme, ticker):
    for tmpl, name in [(t['h5a'], theme), (t['h5b'], theme), (t['h5a'], ticker), (t['h5b'], ticker)]:
        s = fits(tmpl, name, 30)
        if s:
            return s
    raise ValueError(f"H5 unresolved for {theme}/{ticker}")

def resolve_simple(tmpl, theme, ticker, limit):
    for name in (theme, ticker):
        s = fits(tmpl, name, limit)
        if s:
            return s
    raise ValueError(f"unresolved for {theme}/{ticker} tmpl={tmpl}")

def build_row(t, theme, ticker):
    h1, h2, h3, h4 = t['a']
    h5 = resolve_h5(t, theme, ticker)
    h6 = resolve_simple(t['h6'], theme, ticker, 30)
    h7 = resolve_simple(t['h7'], theme, ticker, 30)
    h8 = resolve_simple(t['h8'], theme, ticker, 30)
    h9 = resolve_simple(t['h9'], theme, ticker, 30)
    h10 = resolve_simple(t['h10'], theme, ticker, 30)
    h11, h12, h13, h14, h15 = t['c']
    d1 = t['d1']
    d2 = resolve_simple(t['d2'], theme, ticker, 90)
    d3 = t['d3']
    d4 = t['d4']
    headlines = [h1,h2,h3,h4,h5,h6,h7,h8,h9,h10,h11,h12,h13,h14,h15]
    assert len(headlines) == 15
    assert len(set(headlines)) == 15, f"DUPLICATE headlines: {headlines}"
    for h in headlines:
        assert len(h) <= 30, f"Headline over 30: {h!r} ({len(h)})"
    for d in (d1, d2, d3, d4):
        assert len(d) <= 90, f"Description over 90: {d!r} ({len(d)})"
    for h in headlines + [d1, d2, d3, d4]:
        low = h.lower()
        for bw in BUY_WORDS:
            assert bw not in low, f"Compliance: forbidden word '{bw}' in {h!r}"
    return headlines, [d1, d2, d3, d4]

wb = openpyxl.load_workbook(SRC, data_only=True)
src_ws = wb['English Ads']
headers = [c.value for c in src_ws[1]]
rows = []
for row in src_ws.iter_rows(min_row=2, max_row=src_ws.max_row, values_only=False):
    vals = {headers[i]: row[i].value for i in range(len(headers))}
    if vals['Ad group'] is None:
        continue
    rows.append(vals)
print(f"Loaded {len(rows)} source rows")

bold = Font(bold=True)

for tab_name, t in LANGS.items():
    ws = wb.create_sheet(tab_name)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = bold
    for r_idx, vals in enumerate(rows, start=2):
        theme = vals['Theme']
        ticker = vals['Ticker']
        headlines, descs = build_row(t, theme, ticker)
        path1 = t['path1']
        path2 = vals['Path 2']
        if len(str(path2)) > 15:
            path2 = 'booking-hldgs' if str(path2) == 'booking-holdings' else str(path2)[:15].rstrip('-')
        assert len(path1) <= 15, path1
        assert len(str(path2)) <= 15, f"path2 over limit {path2}"
        out = {
            'Ad group': vals['Ad group'],
            'Theme': vals['Theme'],
            'Main Keyword': t['keyword'].format(name=theme),
            'Ticker': vals['Ticker'],
            'Avg. Monthly Search Volume': None,
            'Final URL': vals['Final URL'],
        }
        for i, h in enumerate(headlines, start=1):
            out[f'Headline {i}'] = h
        out['Description 1'] = descs[0]
        out['Description 1 position'] = '1'
        out['Description 2'] = descs[1]
        out['Description 3'] = descs[2]
        out['Description 4'] = descs[3]
        out['Path 1'] = path1
        out['Path 2'] = path2
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=r_idx, column=c_idx, value=out.get(h))
    ws.freeze_panes = 'A2'
    widths = {'A':20,'B':20,'C':22,'D':10,'E':14,'F':34}
    for i in range(7, len(headers)+1):
        col_letter = ws.cell(row=1, column=i).column_letter
        widths[col_letter] = 26
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    print(f"Built sheet: {tab_name} ({len(rows)} rows)")

wb.save(OUT)
print("Saved:", OUT)
